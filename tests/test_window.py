"""The window end to end, offscreen: the sidebar places, editing, attaching, packages, rentals, exports."""

from __future__ import annotations

from contextlib import contextmanager

import pytest
import shiboken6
from PySide6.QtCore import Qt, QThread, QTimer
from PySide6.QtGui import QContextMenuEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QLabel, QWidget

from bodega.core import sample
from bodega.core.model import RentalLine
from bodega.core.store import Library


@pytest.fixture
def win(qapp, data_folder):
    from bodega.ui.window.main_window import MainWindow

    w = MainWindow(sample.fill(Library(data_folder).load()))
    w.show()
    yield w
    w.close()


def when_open(kind, do) -> None:
    """`do(window)` once a window of this kind shows (it may open from inside a click and wait there)."""
    tries = [500]

    def look() -> None:
        shown = [w for w in QApplication.topLevelWidgets() if isinstance(w, kind) and w.isVisible()]
        if shown:
            do(shown[0])
        elif tries[0]:
            tries[0] -= 1
            QTimer.singleShot(10, look)

    QTimer.singleShot(0, look)


def wait_for(done, ms: int = 5000) -> None:
    for _ in range(ms // 10):
        if done():
            return
        QTest.qWait(10)
    raise AssertionError("never happened")


@contextmanager
def still_there_after(widget: QWidget):
    """A click on `widget` may open a window or a menu; anything that window throws away while it's open is
    gone before the click is over. If that's `widget` itself, Bodega crashes (on a Mac, at once). Fails if it
    happens during the `with`."""
    start = QThread.currentThread().loopLevel()
    thrown = []
    original = QWidget.deleteLater

    def delete_later(w: QWidget) -> None:
        if QThread.currentThread().loopLevel() > start and (w is widget or w.isAncestorOf(widget)):
            thrown.append(type(widget).__name__)
        original(w)

    QWidget.deleteLater = delete_later
    name = type(widget).__name__
    try:
        yield
    finally:
        QWidget.deleteLater = original
    assert shiboken6.isValid(widget), f"the {name} was gone before its own click was over"
    assert not thrown, f"the {thrown[0]} was thrown away in the middle of its own click"


def test_first_launch_is_empty(qapp, data_folder):
    from bodega.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    titles = [t.findChild(QLabel, "StartTitle").text() for t in w.empty.tiles]
    assert w.pages.currentWidget() is w.empty and titles[0] == "Add Manually" and len(titles) == 3
    assert not w.search.isVisible()  # nothing to search yet
    w.go("rentals")
    assert "profile" in w.rentals.empty_title.text().lower()
    assert [t.findChild(QLabel, "StartTitle").text() for t in w.empty.tiles] == ["New Rental"]
    w.close()


def test_empty_gear_page_starts_and_takes_a_dropped_sheet(qapp, data_folder, tmp_path):
    from PySide6.QtCore import QMimeData, QPointF, QUrl
    from PySide6.QtGui import QDropEvent

    from bodega.ui.components.empty import sheet_in
    from bodega.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    w.empty.tiles[0].clicked.emit()  # Add Manually: a new piece, open on the right
    assert len(w.lib.items) == 1 and w.pages.currentWidget() is w.gear_table and w.search.isVisible()
    w.deselect()  # left blank, so it goes
    assert not w.lib.items and w.pages.currentWidget() is w.empty

    def mime(name: str) -> QMimeData:
        m = QMimeData()
        m.setUrls([QUrl.fromLocalFile(str(tmp_path / name))])
        return m

    assert sheet_in(mime("Gear.XLSX")) == tmp_path / "Gear.XLSX"
    assert sheet_in(mime("photo.jpg")) is None

    def drop(name: str) -> list:
        got: list = []
        w.empty.dropped.connect(got.append)
        m = mime(name)  # kept alive: the event only points at it
        e = QDropEvent(
            QPointF(10, 10),
            Qt.DropAction.CopyAction,
            m,
            Qt.MouseButton.LeftButton,
            Qt.KeyboardModifier.NoModifier,
        )
        w.empty.dropEvent(e)
        QApplication.processEvents()
        w.empty.dropped.disconnect(got.append)
        return got

    # (the file isn't there, so the window's import just says it can't read it: no Import window opens)
    assert drop("Gear.csv") == [tmp_path / "Gear.csv"]
    assert drop("photo.jpg") == []
    w.go("kits")  # only the gear page takes spreadsheets
    assert drop("Gear.csv") == []
    w.close()


def test_saving_your_profile_with_the_cursor_in_a_field(qapp, data_folder):
    """Saving in Settings › Profiles puts a fresh form in place of the old one. The old one used to be thrown
    away while it still showed with the focus in it (on a Mac, the field you typed in), and Bodega crashed."""
    from bodega.ui.dialogs.settings import SettingsDialog
    from bodega.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    done = []

    def type_and_save(settings, field: str, text: str, then) -> None:
        f = getattr(settings.profile_editor, field)
        f.setFocus()
        QTest.keyClicks(f, text)
        QTest.mouseClick(settings.profile_editor.save, Qt.MouseButton.LeftButton)
        QTimer.singleShot(100, then)  # Settings stays open after Save, going about its business

    def add(settings) -> None:
        settings.activateWindow()
        QTest.qWaitForWindowActive(settings)  # typing leaves the focus in the field, as it does for you
        type_and_save(settings, "name", "Juniper Lane Films", lambda: change(settings))  # a new profile

    def change(settings) -> None:
        type_and_save(
            settings, "tagline", "Camera · Lighting · Grip", lambda: (done.append(1), settings.close())
        )

    when_open(SettingsDialog, add)
    w.edit_business(True)
    assert done
    b = next(iter(Library(data_folder).load().businesses.values()))
    assert (b.name, b.tagline) == ("Juniper Lane Films", "Camera · Lighting · Grip")
    w.close()


def test_a_new_rental_from_its_tile_before_there_is_a_profile(qapp, data_folder):
    """With no profile yet, the New Rental tile opens Settings › Profiles first. On Coming Up the page swaps its
    tiles meanwhile, and the tile used to be gone before its own click was over."""
    from bodega.ui.dialogs.settings import SettingsDialog
    from bodega.ui.window.main_window import MainWindow

    w = MainWindow(Library(data_folder).load())
    w.show()
    w.go("upcoming")
    tile = w.empty.tiles[0]

    def add(settings) -> None:
        settings.profile_editor.name.setText("Juniper Lane Films")
        settings.profile_editor.save.click()
        settings.close()

    when_open(SettingsDialog, add)
    with still_there_after(tile):
        QTest.mouseClick(tile, Qt.MouseButton.LeftButton)
    wait_for(lambda: w.rentals.r is not None)  # and the first rental is open
    assert [b.name for b in Library(data_folder).load().businesses.values()] == ["Juniper Lane Films"]
    w.close()


def test_renaming_a_group_from_its_row_in_the_sidebar(win, monkeypatch, data_folder):
    """Right-click › Rename Group… takes away the very row that was right-clicked. That used to crash."""
    from PySide6.QtWidgets import QMenu

    from bodega.ui.dialogs import common

    monkeypatch.setattr(common, "ask_name", lambda *a, **k: "Grip & Rigging")
    group = win.lib.groups()[0]
    row = win.sidebar.rows[f"grp:{group}"]

    def rename(menu) -> None:
        next(a for a in menu.actions() if a.text() == "Rename Group…").trigger()
        menu.close()

    when_open(QMenu, rename)
    at = row.rect().center()
    with still_there_after(row):
        QApplication.sendEvent(
            row, QContextMenuEvent(QContextMenuEvent.Reason.Mouse, at, row.mapToGlobal(at))
        )
    wait_for(lambda: "Grip & Rigging" in win.lib.groups())
    assert "grp:Grip & Rigging" in win.sidebar.rows and f"grp:{group}" not in win.sidebar.rows
    assert "Grip & Rigging" in Library(data_folder).load().groups()


def test_a_panel_shows_something_else_from_its_top(win):
    """Scrolled down one rental, then opening another: it starts at its top, not where the last one was."""
    lib = win.lib
    gear = list(lib.items.values())
    a = next(iter(lib.rentals.values()))
    b = lib.new_rental(win.business())
    for r in (a, b):
        r.lines = [RentalLine("item", gear[n % len(gear)].id, 1) for n in range(12)]
        lib.save_rental(r)
    win.go("rentals")
    win.open("rentals", a.id)
    bar = win.rentals.scroll.verticalScrollBar()
    QTest.qWait(10)
    bar.setValue(bar.maximum())
    assert bar.value() > 0
    win.open("rentals", a.id)  # the same one again: it stays where you were
    assert bar.value() == bar.maximum()
    win.open("rentals", b.id)
    assert bar.value() == 0
    win.open("gear", gear[0].id)
    bar = win.items.scroll.verticalScrollBar()
    bar.setValue(bar.maximum())
    win.open("gear", gear[1].id)
    assert bar.value() == 0


def test_a_long_name_shows_all_of_itself(qapp):
    """A name on two lines or more fits its box exactly, so it never scrolls inside it (hiding its first line)
    and the wheel over it scrolls the panel."""
    from PySide6.QtCore import QPoint, QPointF
    from PySide6.QtGui import QWheelEvent

    from bodega.ui.components.widgets import TitleEdit, vbox

    host = QWidget()
    vbox(host, 0, 0).addWidget(t := TitleEdit())
    host.resize(240, 400)
    host.show()
    for words in (1, 6, 30):
        t.setText(" ".join(["Sony FX6 Cinema Camera"] * words))
        QApplication.processEvents()
        assert t.verticalScrollBar().maximum() == 0, t.document().blockCount()
    assert t.document().firstBlock().layout().lineCount() > 2
    e = QWheelEvent(
        QPointF(10, 10), QPointF(10, 10), QPoint(), QPoint(0, -120), Qt.MouseButton.NoButton,
        Qt.KeyboardModifier.NoModifier, Qt.ScrollPhase.NoScrollPhase, False,
    )  # fmt: skip
    QApplication.sendEvent(t.viewport(), e)
    assert not e.isAccepted()  # passed on to the panel
    host.close()


def test_scrolling_past_dropdowns_and_dates_leaves_the_cursor_alone(win):
    """Qt gives dropdowns and dates the focus as the wheel goes over them, taking it from the field you're
    typing in. In Bodega they take it from a click or Tab only."""
    from PySide6.QtWidgets import QComboBox, QDateEdit

    win.go("rentals")
    win.open("rentals", next(iter(win.lib.rentals)))
    QApplication.processEvents()
    controls = [
        w for w in win.findChildren(QWidget) if isinstance(w, (QComboBox, QDateEdit)) and w.isVisible()
    ]
    assert controls
    assert all(w.focusPolicy() == Qt.FocusPolicy.StrongFocus for w in controls)


def test_sidebar_places(win):
    for key in ("gear", "kits", "out", "upcoming", "rentals", "calendar"):
        win.go(key)  # direct: an error inside a signal wouldn't fail the test
        assert win.view_key == key and win.sidebar.rows[key].property("selected")
    win.go("kits")
    assert not win.dock.opened  # nothing picked: no panel
    win.open("packages", next(iter(win.lib.packages)))
    assert win.dock.opened and win.packages.sum_rows["value"].text()
    win.sidebar.pick("gear")
    assert win.pages.currentWidget() is win.gear_table and not win.dock.opened
    win.open("gear", next(iter(win.lib.items)))
    assert win.dock.opened and win.items.it is not None
    win.deselect()  # Esc
    assert not win.dock.opened and win.items.it is None
    grp = win.lib.groups()[0]
    win.sidebar.pick(f"grp:{grp}")
    t = win.gear_table
    ids = [r.data(0, 256) for r in t.item_rows() if not r.isHidden()]
    assert ids and all(win.lib.items[i].group == grp for i in ids)
    assert win.title.text() == grp


def test_page_window(win, monkeypatch):
    from bodega.ui.dialogs.preview import PageDialog

    seen = []
    monkeypatch.setattr(PageDialog, "exec", lambda self: seen.append(self.panel.current_doc()) or 0)
    win.show_pages(("gear", ""), "ratecard")
    r = next(iter(win.lib.rentals.values()))
    win.show_pages(("rental", r.id))
    pkg = next(iter(win.lib.packages.values()))
    win.show_pages(("kit", pkg.id))
    assert seen == ["ratecard", "agreement", "package"]
    win.page_ctx = ("gear", "")
    assert win.view("inventory")[0].kind == "inventory" and win.page_list(doc="inventory")
    win.page_ctx = ("rental", r.id)
    assert win.page_list()


def test_where_and_rent_this(win):
    light = next(i for i in win.lib.items.values() if i.name.startswith("Aputure STORM 80c"))
    win.open("gear", light.id)
    assert win.items.it is light and win.gear_table.current_id() == light.id
    assert win.items.where.isVisibleTo(win.items)  # the sample rental has it booked
    n = len(win.lib.rentals)
    win.rent_this()
    assert len(win.lib.rentals) == n + 1 and win.view_key == "upcoming"
    assert win.rentals.r.lines[0].ref == light.id
    assert win.rental_table.current_id() == win.rentals.r.id


def test_dividers_totals_and_attached(win):
    from bodega.core import pricing

    fx6 = next(i for i in win.lib.items.values() if i.name.startswith("Sony FX6"))
    cage = next(i for i in win.lib.items.values() if i.parent == fx6.id)
    win.go("gear")
    t = win.gear_table
    names = [t.topLevelItem(i).text(1) for i in range(t.topLevelItemCount())]
    assert names == win.lib.groups()  # a divider per group
    shown = [r.data(0, 256) for r in t.item_rows() if not r.isHidden()]
    assert fx6.id in shown and cage.id in shown  # every piece has its own row
    fx6_row = next(r for r in t.item_rows() if r.data(0, 256) == fx6.id)
    assert fx6_row.text(2) == f"+{len(pricing.children(win.lib, fx6.id))}"
    from bodega.ui.window.tables import WITH

    cage_row = next(r for r in t.item_rows() if r.data(0, 256) == cage.id)
    assert cage_row.data(0, WITH) == fx6.name  # says what it goes with
    b = win.lib.pricing()
    everything = pricing.Stock()
    for it in win.lib.items.values():
        everything.add(pricing.stock(win.lib, it, b, attached=False))
    assert win.total_values["value"].text().replace(",", "").endswith(f"{everything.value:.2f}")
    assert everything.pieces == sum(i.qty for i in win.lib.items.values() if i.status != "retired")
    win.search.setText("cage")
    assert win.count.full_text() == "1"
    row = next(r for r in t.item_rows() if not r.isHidden())
    assert row.data(0, 256) == cage.id
    win.open("gear", cage.id)
    assert win.items.it is cage
    win.search.clear()
    win.set_gear_filter("repair")
    assert win.count.full_text() == ""
    win.set_gear_filter("all")
    win.open("gear", fx6.id)
    assert win.items.kids_lay.count() == len(pricing.children(win.lib, fx6.id))


def test_toasts(win):
    import time

    t = win.toaster
    win.toast("Copied 3 pieces.", "done")
    assert t.isVisible() and (t.text, t.tone) == ("Copied 3 pieces.", "done")
    pill = t._pill.translated(t.pos().x(), t.pos().y())
    area = win._toast_area()
    assert area.contains(pill.toRect()) and pill.center().x() == pytest.approx(area.center().x(), abs=2)
    win.toast("Something odd.", "sparkly")  # an unknown tone is plain news
    assert t.tone == ""

    def wait(ms):
        stop = time.monotonic() + ms / 1000
        while time.monotonic() < stop:
            QApplication.processEvents()
            time.sleep(0.01)

    t.dismiss()
    win.toast("Saved.", "done")  # a new one while the last fades away stays up
    wait(500)
    assert t.isVisible() and t.text == "Saved."
    t.dismiss()
    wait(500)
    assert not t.isVisible()


def test_light_and_dark(win):
    from bodega.ui.style import tokens

    win.set_look("light")
    assert tokens.SCHEME == "light" and tokens.LIGHT["BG"] == tokens.BG
    assert win.lib.settings["appearance"] == "light"
    win.set_look("dark")
    assert tokens.DARK["BG"] == tokens.BG


def test_new_item_edit_and_attach(win, monkeypatch):
    from bodega.ui.dialogs import gear_picker

    win.new("gear")
    it = win.items.it
    win.items.name.setText("Matte Box for FX6")
    win.items._set("name", "Matte Box for FX6")
    win.items.qty.setValue(2)
    assert not win.items.blocks["serials"].isVisibleTo(win.items)  # added only when wanted
    assert not win.items.serial_fields  # and its fields are made only then
    win.items._add_part("serials")
    assert win.items.blocks["serials"].isVisibleTo(win.items) and "serials" in it.extras
    assert len(win.items.serial_fields) == 2
    win.items._remove_part("serials")
    assert "serials" not in it.extras
    fx6 = next(i for i in win.lib.items.values() if i.name.startswith("Sony FX6"))

    def pick(self):
        self.picked = [("item", it.id)]
        return gear_picker.GearPicker.DialogCode.Accepted

    monkeypatch.setattr(gear_picker.GearPicker, "exec", pick)
    win.attach_to(fx6.id)
    win.flush()
    again = Library(win.lib.root).load()
    assert again.items[it.id].parent == fx6.id and again.items[it.id].qty == 2
    win.open("gear", it.id)
    assert win.items.blocks["works"].isVisibleTo(win.items)  # it goes with the FX6
    win.detach(it.id)
    assert not win.lib.items[it.id].parent


def test_rental_line_dates_from_the_panel(win):
    win.go("rentals")
    win.open("rentals", next(iter(win.lib.rentals)))
    r = win.rentals.r
    line = RentalLine("item", next(iter(win.lib.items)), 1)
    r.lines.append(line)
    win.rentals.build_lines(open_line=line)
    ed = win.rentals.editors[-1]
    ed.own.setChecked(True)
    assert r.own_dates(line)
    ed.own.setChecked(False)
    assert not line.pickup and not r.own_dates(line)


def test_status_and_numbers(win):
    win.new("rentals")
    r = win.rentals.r
    assert r.status == "quote" and r.number.endswith("-002")
    win.rentals.set_status("booked")
    win.flush()
    assert Library(win.lib.root).load().rentals[r.id].status == "booked"


def test_a_typed_number_another_rental_has_is_said_under_it(win):
    first = next(iter(win.lib.rentals.values()))
    win.new("rentals")
    panel = win.rentals
    panel.number.setText(first.number)
    panel.number.textEdited.emit(first.number)
    assert panel.number_clash.isVisibleTo(panel) and "already has this number" in panel.number_clash.text()
    panel.number.setText(first.number + "b")
    panel.number.textEdited.emit(first.number + "b")
    assert not panel.number_clash.isVisibleTo(panel)


def test_next_number_in_the_profile_form(qapp, data_folder):
    from bodega.ui.dialogs.profile import ProfileDialog

    lib = sample.fill(Library(data_folder).load())
    b = next(iter(lib.businesses.values()))
    d = ProfileDialog(b, False, next_for=lambda prefix: lib.number_after(prefix, 2026))
    free = lib.number_after(b.prefix, 2026)
    assert d.next_field.text() == f"{free:03d}" and not d.save.isEnabled()  # nothing changed yet
    d.next_field.setText("000")
    d.next_field.textEdited.emit("000")
    assert not d.save.isEnabled() and not d.next_note.isHidden()
    d.next_field.setText("120")
    d.next_field.textEdited.emit("120")
    assert d.save.isEnabled() and d.next_note.isHidden()
    d._save()
    assert d.next_number == 120


def test_import_into_the_window(win, tmp_path):
    import openpyxl

    from tests.test_importer import ROWS

    wb = openpyxl.Workbook()
    for row in ROWS:
        wb.active.append(row)
    path = tmp_path / "list.xlsx"
    wb.save(path)
    before = len(win.lib.items)
    assert win.import_file(path, ask=False) == 9  # the sample already has the FX6's top handle
    assert len(win.lib.items) == before + 9
    assert win.import_file(path, ask=False) == 0  # nothing twice


def test_export_pdf(win, tmp_path):
    from bodega.doc import export

    win.page_ctx = ("rental", next(iter(win.lib.rentals)))
    out = export.to_pdf(win.page_list(), tmp_path / "a.pdf")
    assert out.stat().st_size > 1000


def test_attach_and_add_through_the_picker(win, monkeypatch):
    from bodega.ui.dialogs import gear_picker

    picks: list = []

    def fake_exec(self):
        self.picked = list(picks)
        return gear_picker.GearPicker.DialogCode.Accepted

    monkeypatch.setattr(gear_picker.GearPicker, "exec", fake_exec)
    lib = win.lib
    mic = next(i for i in lib.items.values() if i.name == "Shotgun Microphone")
    pole = next(i for i in lib.items.values() if i.name.startswith("Boompole"))
    picks[:] = [("item", pole.id)]
    win.attach_to(mic.id)
    assert pole.parent == mic.id
    pkg = next(iter(lib.packages.values()))
    n = len(pkg.lines)
    picks[:] = [("item", mic.id)]
    win.add_to_package(pkg.id)
    assert len(pkg.lines) == n + 1
    win.go("rentals")
    win.open("rentals", next(iter(win.lib.rentals)))
    r = win.rentals.r
    n = len(r.lines)
    picks[:] = [("package", pkg.id), ("item", mic.id)]
    win.add_to_rental(r.id)
    assert len(r.lines) == n + 1  # the package was already on it: one more of it
    assert r.lines[0].qty == 2


def test_delete_with_confirm(win, monkeypatch):
    from bodega.ui.window import editing

    asked = []
    monkeypatch.setattr(editing, "confirm", lambda _w, title, *a, **k: asked.append(title) or True)
    fx6 = next(i for i in win.lib.items.values() if i.name.startswith("Sony FX6"))
    assert win.lib.rentals_with(fx6.id)
    win.delete("gear", fx6.id)  # on a rental: retired, never deleted
    assert asked[-1] == f"Retire {fx6.name} Instead?" and win.lib.items[fx6.id].status == "retired"
    free = next(i for i in win.lib.items.values() if not win.lib.rentals_with(i.id) and not i.parent)
    win.delete("gear", free.id)
    assert free.id not in win.lib.items
    win.go("rentals")
    win.open("rentals", next(iter(win.lib.rentals)))
    r = win.rentals.r
    assert r.status == "booked"
    win.delete("rentals", r.id)  # booked: cancelled and kept
    assert asked[-1] == f"Mark {r.number} as Cancelled?" and win.lib.rentals[r.id].status == "cancelled"
    assert [e.what for e in Library(win.lib.root).load().rentals[r.id].history][-1] == "Cancelled"
    r.status = "quote"
    win.rentals.load(r)
    win.delete("rentals", r.id)  # a quote can go
    assert not win.lib.rentals
    assert win.pages.currentWidget() is win.empty


def test_import_window_choices(qapp, data_folder, tmp_path):
    """The Import window: departments and single pieces can be left out, columns re-picked."""
    import openpyxl

    from bodega.ui.dialogs.import_sheet import ImportDialog
    from tests.test_importer import ROWS

    wb = openpyxl.Workbook()
    for row in ROWS:
        wb.active.append(row)
    path = tmp_path / "list.xlsx"
    wb.save(path)
    from bodega.core import importer

    d = ImportDialog(path, importer.load_rows(path), set())
    assert d.go.isEnabled() and len(d.chosen()) == 10
    d.dep_toggles["G&E"].setChecked(False)
    assert all(i.department != "G&E" for i in d.chosen())
    d.list.item(0).setCheckState(__import__("PySide6.QtCore", fromlist=["Qt"]).Qt.CheckState.Unchecked)
    assert len(d.chosen()) == 5
    d.attach.setChecked(False)
    assert not any(i.parent for i in d.chosen())
    assert len(d.chosen()) == 5  # what was unticked stays unticked after a re-read
    d.pick_column("name", None)
    assert not d.chosen() and not d.go.isEnabled()
    d.pick_column("name", 1)
    d._done()
    assert d.items and d.pricing == ("4", "8.25")  # no profile yet: it starts with the sheet's pricing


def test_an_import_never_replaces_your_pricing_unasked_and_says_what_it_filled_in(
    qapp, data_folder, tmp_path
):
    from bodega.core import importer
    from bodega.ui.dialogs.import_sheet import ImportDialog

    sheet = tmp_path / "list.csv"
    sheet.write_text("Name,Qty,Price\nAputure STORM 80c,,500\nAputure STORM 1200x,2,2000\n")
    d = ImportDialog(sheet, importer.load_rows(sheet), set(), has_profile=True)
    assert not d.use_pricing.isChecked()  # you have a profile: its pricing stays yours
    assert (
        "1 with no quantity count as 1" in d.filled.text()
        and "2 with no Rent? say are for rent" in d.filled.text()
    )
    assert "company" in d.filled.text() and all(i.brand == "Aputure" for i in d.chosen())
    d.guess.setChecked(False)
    assert not any(i.brand for i in d.chosen()) and "company" not in d.filled.text()


def test_the_sheets_links_go_to_gear_you_have_only_if_you_say(win, tmp_path, monkeypatch):
    from bodega.ui.dialogs import import_sheet

    mic = next(i for i in win.lib.items.values() if i.name == "Shotgun Microphone")
    mic.link = ""
    sheet = tmp_path / "list.csv"
    sheet.write_text(f"Name,Department,Link\n{mic.name},{mic.department},https://maker.example/mic\n")

    def run(links_on: bool):
        def exec_(d):
            assert not d.links.isHidden()
            d.links.setChecked(links_on)
            d._done()
            return import_sheet.ImportDialog.DialogCode.Accepted

        monkeypatch.setattr(import_sheet.ImportDialog, "exec", exec_)
        win.import_file(sheet)

    run(False)
    assert not win.lib.items[mic.id].link
    run(True)
    assert win.lib.items[mic.id].link == "https://maker.example/mic"


def test_your_own_groups(win, monkeypatch):
    from bodega.ui.dialogs import common as dialogs

    lib = win.lib
    mic = next(i for i in lib.items.values() if i.name == "Shotgun Microphone")
    pole = next(i for i in lib.items.values() if i.name.startswith("Boompole"))
    monkeypatch.setattr(dialogs, "ask_name", lambda *a, **k: "Cases")
    win.new_group()
    assert "Cases" in lib.groups() and win.view_key == "grp:Cases"
    win.go("gear")
    win.move_gear([mic.id, pole.id], "Cases")
    assert mic.group == pole.group == "Cases"
    t = win.gear_table
    assert "Cases" in [t.topLevelItem(i).text(1) for i in range(t.topLevelItemCount())]
    t.dropped.emit([mic.id], "Lights")  # dragged onto the Lights divider
    assert mic.group == "Lights"
    monkeypatch.setattr(dialogs, "ask_name", lambda *a, **k: "Lighting")
    win.rename_group("Lights")
    assert mic.group == "Lighting" and "Lights" not in lib.groups()
    win._step_group("Lighting", -10)
    assert lib.groups()[0] == "Lighting"
    assert win.sidebar.rows["grp:Lighting"]
    # the inspector picks from the list, or names a new one
    win.open("gear", mic.id)
    c = win.items.group
    assert c.itemText(c.count() - 1) == "New Group…"
    monkeypatch.setattr(dialogs, "ask_name", lambda *a, **k: "Shotguns")
    c.setCurrentIndex(c.count() - 1)
    win.items._group_picked()
    assert mic.group == "Shotguns" and "Shotguns" in lib.groups()
    again = Library(lib.root).load()
    assert again.items[mic.id].group == "Shotguns"


def test_tidy_up(win, monkeypatch):
    from bodega.ui.dialogs import common as dialogs

    monkeypatch.setattr(dialogs, "tidy_preview", lambda *a, **k: True)
    card = next(i for i in win.lib.items.values() if "CFexpress" in i.name)
    win.tidy_groups()
    assert card.group == "Media" and "grp:Media" in win.sidebar.rows


def test_select_all_and_delete(win, monkeypatch):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QKeyEvent, QKeySequence

    from bodega.ui.window import editing

    asked = []
    monkeypatch.setattr(editing, "confirm", lambda *a, **k: asked.append(a[1]) or True)
    win.go("grp:Microphones")
    t = win.gear_table
    t.setFocus()
    t.select_showing()
    ids = t.chosen_ids()
    audio = [i.id for i in win.lib.items.values() if i.group == "Microphones"]
    assert sorted(ids) == sorted(audio) and len(ids) > 1
    assert "selected" in win.count.full_text()
    before = len(win.lib.items)
    t.keyPressEvent(QKeyEvent(QKeyEvent.Type.KeyPress, Qt.Key.Key_Delete, Qt.KeyboardModifier.NoModifier))
    assert asked and asked[-1].startswith(f"Delete {len(ids)} pieces")
    assert len(win.lib.items) == before - len(ids)
    assert all(i not in Library(win.lib.root).load().items for i in ids)
    # Cmd+A then Delete on everything
    win.go("gear")
    t.select_showing()
    win.delete_gear(t.chosen_ids())
    assert (
        asked[-1].startswith("Retire ") and " and Delete " in asked[-1]
    )  # one alert: retire some, delete the rest
    left = list(win.lib.items.values())
    assert left and all(i.status == "retired" for i in left)  # with what's attached to them
    assert all(win.lib.rentals_with(i.id) or i.parent in win.lib.items for i in left)
    assert QKeySequence.StandardKey.SelectAll


def test_order_columns_and_divide(win):
    from bodega.ui.window.tables import GEAR_COL

    lib = win.lib
    win.go("grp:Body")
    win.view_menu.divide = "none"
    win._view_picked()
    t = win.gear_table
    shown = [r.data(0, 256) for r in t.showing_rows()]
    first, last = shown[0], shown[-1]
    win._reordered([last], first, False)  # dragged by its grip to the top
    assert next(r.data(0, 256) for r in t.showing_rows()) == last
    assert Library(lib.root).load().ordered_items()[0].id == last
    # divide by company
    win.view_menu.divide = "brand"
    win._view_picked()
    names = [t.topLevelItem(i).text(1) for i in range(t.topLevelItemCount())]
    assert "Sony" in names and t.isColumnHidden(GEAR_COL["brand"])
    assert win._divide_for("grp:Body") == "brand" and win._divide_for("gear") == "group"
    # columns you choose
    t._toggle_column("series", True)
    assert not t.isColumnHidden(GEAR_COL["series"]) and "series" not in lib.settings["hidden_columns"]
    # duplicate from the ⋯ menu
    n = len(lib.items)
    win.open("gear", first)
    win.items.duplicate_wanted.emit(first)
    assert len(lib.items) == n + 1


def test_moves_wait_for_the_drag(win, monkeypatch):
    """A drop arrives while the drag is still running; the table must not refill until it's over."""
    from PySide6.QtWidgets import QTreeWidget

    t = win.gear_table
    seen = []

    def fake_drag(self, actions):
        self.after_drag(lambda: seen.append(self.topLevelItemCount()))
        assert not seen  # waits
        self._fill_start()  # anything refilling mid-drag would show here; the queued move hasn't run

    monkeypatch.setattr(QTreeWidget, "startDrag", fake_drag)
    t.startDrag(None)
    assert seen and not t._dragging


def test_sidebar_nests_groups_and_kits(win):
    sb = win.sidebar
    assert not sb.boxes["gear"].isVisible() and sb.rows["gear"].folded  # groups tucked away at first
    sb._fold("gear", False)
    assert sb.boxes["gear"].isVisible() and "gear" not in win.lib.settings["sidebar_folded"]
    kit = next(iter(win.lib.packages.values()))
    win._place_picked(f"kit:{kit.id}")
    assert win.view_key == "kits" and win.packages.pkg is kit
    assert sb.rows[f"kit:{kit.id}"].property("selected")


def test_settings_window(win):
    from bodega.ui.dialogs.settings import SettingsDialog

    d = SettingsDialog(win.lib, win)
    d.look.changed.emit("light")
    assert win.lib.settings["appearance"] == "light"
    d.auto.setChecked(False)
    assert win.lib.settings["auto_update"] is False
    d.close()


def test_back_up_zips_the_data(tmp_path):
    import zipfile

    from bodega.core.fileutil import zip_folder

    src = tmp_path / "Bodega"
    (src / "rentals").mkdir(parents=True)
    (src / "gear.json").write_text("{}")
    (src / "rentals" / "a.json").write_text("{}")
    (src / "ui").mkdir()
    (src / "ui" / "x.svg").write_text("<svg/>")
    dest = tmp_path / "backup.zip"
    assert zip_folder(src, dest, skip=("ui",)) == 2
    assert sorted(zipfile.ZipFile(dest).namelist()) == ["Bodega/gear.json", "Bodega/rentals/a.json"]


def test_dragging_a_row_shifts_the_others_live_and_moves_it(win):
    from PySide6.QtWidgets import QApplication

    from bodega.ui.window.tables import ID

    t = win.gear_table
    win.go("gear")
    rows = [r for r in t.item_rows() if not r.isHidden()]
    first, third = rows[0], rows[2]
    assert first.parent() is third.parent()  # the sample's first group has room to drag in
    first_id, third_id = first.data(0, ID), third.data(0, ID)  # the rows are rebuilt after the drop
    y0 = t.visualItemRect(first).center().y()
    t.start_lift(first, y0)
    lift = t._lift
    assert lift is not None and lift.ids == [first_id]
    t._lift.move_to(t.visualItemRect(third).bottom() + 2)  # past the third row
    for _ in range(30):  # let the rows slide
        lift._step()
    QApplication.processEvents()
    assert lift.k > lift.start_k
    moved = []
    t.reordered.connect(lambda ids, anchor, after: moved.append((ids, anchor, after)))
    t.end_lift(drop=True)
    QApplication.processEvents()
    assert moved and moved[0][0] == [first_id]
    order = list(win.lib.items)
    assert order.index(first_id) > order.index(third_id)


def test_esc_puts_a_lifted_row_back(win):
    t = win.gear_table
    win.go("gear")
    row = next(r for r in t.item_rows() if not r.isHidden())
    before = list(win.lib.items)
    t.start_lift(row, t.visualItemRect(row).center().y())
    t._lift.move_to(t.viewport().height() - 5)
    t.end_lift(drop=False)
    assert t._lift is None and list(win.lib.items) == before


def test_all_rentals_starts_folded_even_with_older_settings(qapp, data_folder):
    from bodega.ui.window.main_window import MainWindow

    lib = sample.fill(Library(data_folder).load())
    lib.settings["sidebar_folded"] = ["gear"]  # saved before All Rentals existed
    w = MainWindow(lib)
    assert w.sidebar.rows["rentals"].folded and not w.sidebar.boxes["rentals"].isVisibleTo(w.sidebar)
    w.sidebar._fold("rentals", False)  # you open it: that's remembered
    w.close()
    w2 = MainWindow(Library(data_folder).load())
    assert not w2.sidebar.rows["rentals"].folded
    w2.close()


def test_calendar_month_and_timeline(win):
    from PySide6.QtCore import QPoint
    from PySide6.QtTest import QTest

    from bodega.core import dates, pricing

    win.go("calendar")
    cal = win.calendar
    cal.resize(1200, 800)
    cal.set_mode("month")
    assert cal.title.text() == f"{dates.MONTH_NAMES[dates.today().month - 1]} {dates.today().year}"
    r = next(iter(win.lib.rentals.values()))
    a = next(a for a, _z in pricing.periods(r) if a is not None)
    cal.month.year, cal.month.month = a.year, a.month
    cal.month.fill(win.lib)
    cal.month.grab()  # paints: lays out the bars
    assert cal.month._bars, "the rental shows as a bar"
    got = []
    cal.picked.connect(got.append)
    c = cal.month._bars[0][0].center()
    QTest.mouseClick(cal.month, Qt.MouseButton.LeftButton, pos=QPoint(int(c.x()), int(c.y())))
    assert got == [r.id]  # a click on a bar opens the rental
    cal.month.step(1)
    cal.month.step(-13)
    assert (cal.month.year, cal.month.month) == (a.year - 1, a.month)
    cal.set_mode("timeline")
    assert win.lib.settings["calendar_mode"] == "timeline" and "–" in cal.title.text()
    start = cal.timeline.start
    cal._step(1)
    assert cal.timeline.start == dates.plus(start, cal.timeline.days)
    cal._today()
    cal.set_mode("month")


def test_filter_button(win):
    win.go("gear")
    f = win.view_menu
    assert not f.property("active") and f.is_default()
    win.set_gear_filter("repair")
    assert f.property("active") and "In Repair" in f.toolTip()
    win.set_gear_filter("all")
    sony = sum(1 for i in win.lib.items.values() if i.brand == "Sony")
    f.filter.company = "Sony"
    f.changed.emit()
    assert win.count.full_text() == str(sony)
    f.filter.missing = {"link"}
    win.refresh()
    shown = [r.data(0, 256) for r in win.gear_table.showing_rows()]
    assert shown and all(not win.lib.items[i].link and win.lib.items[i].brand == "Sony" for i in shown)
    f.m.aboutToShow.emit()  # the menu builds, with Divide By first
    texts = [a.text() for a in f.m.actions() if a.text()]
    assert texts[:3] == ["Group", "Company", "No Dividers"] and "Clear Filters" in texts
    clear = next(a for a in f.m.actions() if a.text() == "Clear Filters")
    assert clear.isEnabled()
    clear.trigger()
    assert f.is_default() and not f.property("active")
    win.go("kits")
    assert f.isHidden()


def test_undo_and_redo(win, monkeypatch):
    from bodega.ui.window import editing

    monkeypatch.setattr(editing, "confirm", lambda *a, **k: True)
    lib = win.lib
    n = len(lib.items)
    victim = next(i for i in lib.items if not lib.rentals_with(i))  # (gear on a rental is retired instead)
    name = lib.items[victim].name
    win.gear_table.setFocus()
    win.delete_gear([victim])
    assert victim not in lib.items
    win.undo()
    assert victim in lib.items and lib.items[victim].name == name and len(lib.items) == n
    assert victim in Library(lib.root).load().items  # saved back
    win.redo()
    assert victim not in lib.items and victim not in Library(lib.root).load().items
    win.undo()
    # a change in a panel, then undo with the panel open: the panel shows the old name
    win.open("gear", victim)
    win.items.it.name = "Renamed"
    win._item_edited(win.items.it)
    win.flush()
    win.setFocus()
    win.undo()
    assert lib.items[victim].name == name and win.items.it is lib.items[victim]
    # groups too
    win.lib.add_group("Undo Test")
    win.refresh()
    win.undo()
    assert "Undo Test" not in win.lib.groups()


def test_blank_new_piece_goes_on_esc(win):
    n = len(win.lib.items)
    win.new("gear")
    assert len(win.lib.items) == n + 1
    steps = len(win.history.undos)
    win.deselect()  # Esc
    assert len(win.lib.items) == n and len(Library(win.lib.root).load().items) == n
    assert len(win.history.undos) == steps - 1  # and it left no undo step
    win.new("gear")
    win.items.it.name = "Kept"
    win._item_edited(win.items.it)
    win.deselect()
    assert any(i.name == "Kept" for i in win.lib.items.values())
    k = len(win.lib.packages)
    win.new("packages")
    win.deselect()
    assert len(win.lib.packages) == k


def test_alert_return_says_yes(win):
    from PySide6.QtCore import QTimer
    from PySide6.QtTest import QTest

    from bodega.ui.dialogs.common import confirm

    QTimer.singleShot(50, lambda: QTest.keyClick(QApplication.activeModalWidget(), Qt.Key.Key_Return))
    assert confirm(win, "Delete it?", "", "Delete") is True
    QTimer.singleShot(50, lambda: QTest.keyClick(QApplication.activeModalWidget(), Qt.Key.Key_Escape))
    assert confirm(win, "Delete it?", "", "Delete") is False


def test_shortcuts_listed(win):
    from bodega.ui.dialogs.shortcuts import keys_by_menu

    menus = dict(keys_by_menu(win.menuBar()))
    edit = dict(menus["Edit"])
    assert {"Undo", "Redo", "Copy", "Select All", "Find"} <= set(edit)
    assert "New" in dict(menus["File"])


def test_drifting_left_or_up_keeps_the_row_lifted(win):
    from PySide6.QtCore import QEvent, QPointF
    from PySide6.QtGui import QMouseEvent

    t = win.gear_table
    win.go("gear")
    row = [r for r in t.item_rows() if not r.isHidden()][1]
    y = t.visualItemRect(row).center().y()
    t.start_lift(row, y)
    vp = t.viewport()
    for x, dy in ((4, -6), (-6, -14), (1, -40)):  # back up, and a little past the grip's edge
        local = QPointF(x, y + dy)
        ev = QMouseEvent(
            QEvent.Type.MouseMove,
            local,
            QPointF(vp.mapToGlobal(local.toPoint())),
            Qt.MouseButton.NoButton,
            Qt.MouseButton.LeftButton,
            Qt.KeyboardModifier.NoModifier,
        )
        t.mouseMoveEvent(ev)
        assert t._lift is not None, (x, dy)  # still following the pointer
    win.refresh()  # a refill asked for mid-drag waits until the row lands
    assert t._lift is not None and t._fill_later is not None
    t.end_lift(drop=False)
    QApplication.processEvents()
    assert t._fill_later is None


def test_ctrl_backspace_while_typing_edits_the_field_and_deletes_nothing(win, monkeypatch):
    """⌘⌫ / Ctrl+Backspace is Delete… for the table, and the editing key in a text field (a word, or the line)."""
    from bodega.ui.window import editing

    asked = []
    monkeypatch.setattr(editing, "confirm", lambda *a, **k: asked.append(a) or False)
    it = next(iter(win.lib.items.values()))
    win.open("gear", it.id)
    win.activateWindow()
    QTest.qWaitForWindowActive(win)
    from PySide6.QtWidgets import QLineEdit

    field = next(f for f in win.items.findChildren(QLineEdit) if f.isVisible() and not f.isReadOnly())
    field.setFocus()
    field.setText("Aputure Amaran")
    field.end(False)
    QTest.keyClick(field, Qt.Key.Key_Backspace, Qt.KeyboardModifier.ControlModifier)
    assert not asked and it.id in win.lib.items  # no "Delete …?" while typing
    assert field.text() != "Aputure Amaran"  # the field took the key


def test_screen_readers_hear_a_name_for_every_icon_button_and_place(win):
    from PySide6.QtWidgets import QAbstractButton

    from bodega.ui.components.panels import NavRow

    QApplication.processEvents()
    unnamed = [
        b.toolTip() or b.objectName()
        for b in win.findChildren(QAbstractButton)
        if b.isVisible() and not b.text().strip() and not b.accessibleName() and b.toolTip()
    ]
    assert not unnamed
    assert win.btn_ai.accessibleName() == "AI Assistant"
    rows = [r for r in win.sidebar.findChildren(NavRow) if r.isVisible()]
    assert rows and all(r.accessibleName() for r in rows)


def test_words_cut_short_can_still_be_read_whole(win):
    """A gear name too long for its column shows whole on hover; a cut label is read out whole."""
    from PySide6.QtWidgets import QStyleOptionViewItem

    from bodega.ui.components.widgets import ElidedLabel
    from bodega.ui.window.tables import GEAR_COL, RowPainter

    pieces = list(win.lib.items.values())
    long, short = pieces[0], min(pieces[1:], key=lambda i: len(i.name))
    long.name = "Sony FX6 Full-Frame Cinema Camera, the Long Name a Rental House Gives It, Twice Over " * 2
    long.parent = ""
    win.refresh()
    QApplication.processEvents()
    t = win.gear_table
    rows = {r.data(0, 256): r for r in t.item_rows()}
    option = QStyleOptionViewItem()
    option.fontMetrics = t.fontMetrics()

    def tip(it):
        return RowPainter.cut_short(t, option, t.indexFromItem(rows[it.id], GEAR_COL["name"]))

    assert tip(long) == long.name and tip(short) == ""
    label = ElidedLabel("A Very Long Label That Has to Be Cut Short")
    label.resize(60, 20)
    label.show()
    QApplication.processEvents()
    assert label.text() != label.full_text() and label.accessibleName() == label.full_text()
    label.deleteLater()


def test_a_search_that_finds_nothing_says_so_with_a_way_back(win):
    win.go("all")
    win.search.setText("zzqx nothing like it")
    QApplication.processEvents()
    nr = win.no_results
    assert nr.isVisibleTo(win) and "zzqx" in nr.detail.text() and nr.clear.text() == "Clear Search"
    nr.clear.click()
    QApplication.processEvents()
    assert win.search.text() == "" and not nr.isVisibleTo(win)
    win.search.setText("cafe zyxw")  # accents never stop a match
    it = next(iter(win.lib.items.values()))
    it.name = "Café Zyxw Lens Kit"
    win.refresh()
    QApplication.processEvents()
    rows = {r.data(0, 256): r for r in win.gear_table.item_rows()}
    assert not nr.isVisibleTo(win) and not rows[it.id].isHidden()  # (with what goes with it)


def test_measured_to_the_standard(win):
    """The standard's Appendix A, measured in the real window: controls 28, the bar's circles 36, table rows 44,
    sidebar rows 32, figures that line up."""
    from PySide6.QtGui import QFont
    from PySide6.QtWidgets import QComboBox, QLineEdit, QPushButton

    from bodega.ui.components.panels import NavRow
    from bodega.ui.style import tokens

    it = next(iter(win.lib.items.values()))
    win.open("gear", it.id)
    QApplication.processEvents()
    controls = [
        w
        for w in win.items.findChildren(QWidget)
        if isinstance(w, (QPushButton, QLineEdit, QComboBox))
        and w.isVisible()
        and w.objectName() not in ("LinkBtn", "Disclosure", "Where", "AddChip")
    ]
    assert controls and {w.height() for w in controls} == {tokens.CONTROL_H}
    for b in (win.btn_new, win.btn_ai, win.btn_export):
        assert (b.width(), b.height()) == (tokens.BAR_CTRL_H, tokens.BAR_CTRL_H)
    assert win.search.height() == tokens.BAR_CTRL_H
    row = win.gear_table.visualItemRect(win.gear_table.topLevelItem(0))
    assert row.height() == tokens.ROW_H
    assert {r.height() for r in win.sidebar.findChildren(NavRow) if r.isVisible()} == {tokens.CHOICE_ROW_H}
    from bodega.ui.window.tables import GEAR_COL, ID

    piece = next(r for r in win.gear_table.item_rows() if r.data(0, ID))
    assert piece.font(GEAR_COL["qty"]).isFeatureSet(QFont.Tag("tnum"))  # figures line up…
    assert not piece.font(GEAR_COL["name"]).isFeatureSet(QFont.Tag("tnum"))  # …words keep their spacing


def test_rental_columns_make_room_for_the_renter(win):
    from bodega.ui.style import tokens
    from bodega.ui.window.tables import RentalTable

    t = RentalTable()
    t.fill(win.lib, "all")
    t.resize(1300, 400)
    t.show()
    QApplication.processEvents()
    assert not any(t.isColumnHidden(c) for c in range(6))
    t.resize(640, 400)  # a 1280 window with a rental's panel open
    QApplication.processEvents()
    assert t.isColumnHidden(2) and t.isColumnHidden(5)  # Production steps aside first, then Total
    assert t.columnWidth(1) >= tokens.RENTER_MIN - 1
    assert not any(t.isColumnHidden(c) for c in (0, 1, 3, 4))  # number, renter, dates and status stay
    t.close()
    t.deleteLater()


def test_a_date_is_never_cut_off(qapp):
    from bodega.ui.components.widgets import DATE_FORMATS, date_field

    d = date_field()
    d.show()
    d.resize(320, 28)
    QApplication.processEvents()
    assert d.displayFormat() == DATE_FORMATS[0]  # "Sun, Oct 18, 2026" when there's room
    d.resize(140, 28)
    QApplication.processEvents()
    assert d.displayFormat() != DATE_FORMATS[0]
    assert d.fontMetrics().horizontalAdvance(d.text()) <= d.width()
    d.close()
    d.deleteLater()


def test_another_place_starts_unsearched_and_unfiltered(win):
    win.go("gear")
    win.search.setText("sony")
    win.view_menu.filter.show = "out"
    win.go("rentals")
    assert win.search.text() == "" and not win.view_menu.filter.active()
    win.go("gear")
    win.search.setText("sony")
    win.go("gear")  # the same place again (a refresh) keeps what you typed
    assert win.search.text() == "sony"


def test_late_says_so_in_words_not_only_in_red(win):
    from datetime import timedelta

    from bodega.core import dates

    r = next(iter(win.lib.rentals.values()))
    r.status = "out"
    r.pickup = dates.iso(dates.today() - timedelta(days=5))
    r.return_date = dates.iso(dates.today() - timedelta(days=2))
    for line in r.lines:
        line.pickup = line.return_date = ""  # every line on the rental's dates
    win.lib.save_rental(r)
    win.refresh()
    assert win.sidebar.rows["out"].count.text() == "1 late"


def test_a_toast_never_hides_the_row_youre_on(win):
    from PySide6.QtCore import QRect
    from PySide6.QtWidgets import QAbstractItemView

    win.go("gear")
    t = win.gear_table
    rows = [r for r in t.showing_rows()]
    row = rows[len(rows) // 2]
    t.setCurrentItem(row)
    t.scrollToItem(row, QAbstractItemView.ScrollHint.PositionAtBottom)
    QApplication.processEvents()
    win.toast("Saved Sony FX6.", "done")
    QApplication.processEvents()

    def on_window():
        rect = t.visualItemRect(row)
        return QRect(t.viewport().mapTo(win, rect.topLeft()), rect.size())

    assert not on_window().intersects(win.toaster.geometry())


def test_from_the_search_the_keyboard_goes_on_into_the_list(win):
    win.go("gear")
    win.search.setFocus()
    QTest.keyClicks(win.search, "sony")
    QTest.keyClick(win.search, Qt.Key.Key_Return)  # Return on words opens the first match
    QApplication.processEvents()
    first = next(iter(win.gear_table.showing_rows()))
    assert win.gear_table.currentItem() is first and win.items.it is not None
    assert win.items.it.id == first.data(0, 256)
    win.search.setFocus()
    QTest.keyClick(win.search, Qt.Key.Key_Down)
    assert win.gear_table.hasFocus()


def test_back_up_runs_without_freezing_the_window(win, monkeypatch, tmp_path):
    from PySide6.QtWidgets import QFileDialog

    target = tmp_path / "Backup.zip"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_a, **_k: (str(target), ""))
    said = []
    monkeypatch.setattr(win, "toast", lambda text, tone="": said.append(text))
    win.back_up()
    assert said == ["Backing up…"]  # the window carries on while the zip is made
    wait_for(lambda: len(said) == 2)
    assert target.exists() and said[1].startswith("Backed up.")


def test_the_totals_say_how_many_pieces_have_no_price(win):
    win.go("gear")
    assert win.total_labels["paid"].text() == "Paid" or "not priced" in win.total_labels["paid"].text()
    it = next(i for i in win.lib.items.values() if i.value)
    it.value = ""
    win.refresh()
    assert "not priced" in win.total_labels["paid"].text()


def test_the_places_come_from_one_list(win):
    from bodega.ui.window.common import PLACES

    for p in PLACES:
        assert win.sidebar.rows[p.key].label.text() == p.title
    view = next(m for m in win.menuBar().findChildren(type(win.view_menu.menu())) if m.title() == "View")
    shown = {a.text(): a.shortcut().toString() for a in view.actions()}
    for p in PLACES:
        if p.keys:
            assert shown[p.menu or p.title] == p.keys


def test_gear_moves_up_and_down_from_the_keyboard(win):
    win.go("gear")
    t = win.gear_table
    rows = [r for r in t.showing_rows()]
    first, second = rows[0], rows[1]
    assert first.parent() is second.parent()
    a, b = first.data(0, 256), second.data(0, 256)
    win.step_gear(1, [a])
    order = [i.id for i in win.lib.ordered_items()]
    assert order.index(a) > order.index(b)
    said = []
    win.toast = lambda text, tone="": said.append(text)
    win.step_gear(-1, [b])  # b is first in its group now
    assert said and "first in its group" in said[0]


def test_export_csv_exports_what_the_list_shows(win, monkeypatch, tmp_path):
    import csv

    from PySide6.QtWidgets import QFileDialog

    target = tmp_path / "gear.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_a, **_k: (str(target), ""))
    win.go("gear")
    win.export_csv()
    every = len(list(csv.reader(target.open(encoding="utf-8-sig"))))
    win.search.setText("sony")
    win.export_csv()
    rows = list(csv.reader(target.open(encoding="utf-8-sig")))
    assert 1 < len(rows) < every and all("sony" in " ".join(r).casefold() for r in rows[1:] if r)


def test_folding_by_keyboard_is_remembered_and_option_click_folds_every_group(win, monkeypatch):
    from bodega.ui.window.tables import GearTable

    win.go("gear")
    t = win.gear_table
    dividers = t._dividers()
    assert len(dividers) > 1 and all(d.isExpanded() for d in dividers)
    t.setCurrentItem(dividers[0])
    t.setFocus()
    QTest.keyClick(t, Qt.Key.Key_Left)
    assert not dividers[0].isExpanded() and win.lib.settings["folded"]  # kept, as a click would be
    monkeypatch.setattr(
        QApplication, "keyboardModifiers", staticmethod(lambda: Qt.KeyboardModifier.AltModifier)
    )
    GearTable._divider_clicked(t, dividers[1], 0)  # ⌥-click on an open one folds them all
    assert not any(d.isExpanded() for d in t._dividers())
