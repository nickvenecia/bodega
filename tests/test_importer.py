"""Spreadsheets in and out (core/importer.py)."""

from __future__ import annotations

import pytest

from bodega.core import importer
from bodega.core.model import Item

# A hand-made list laid out the way people keep them: a title, a header, headings in the first column,
# group headings in the name column, formulas and a total row.
ROWS = [
    ["EQUIPMENT LIST", None, None, None, None, None],
    ["Qty", "Equipment", "Replacement Cost (Per/Unit)", "Total", "Rent?", 0.04],
    ["Camera", None, None, None, None, None],
    ["Body", None, None, None, None, None],
    [1, "Sony FX6 Full-Frame Cinema Camera (Body Only)", 5998, "=(C5*1.0825)*A5", True, "=IF(...)"],
    [1, "Tilta Camera Cage for Sony FX6 with V-Mount Battery Plate", 279, "=(C6*1.0825)*A6", True, None],
    [1, "Tilta Top Handle for Sony FX6", 69, None, True, None],
    [1, "Kondor Blue XLR Handle Extension for FX6 Top Handle", 69, None, True, None],
    ["Batteries", None, None, None, None, None],
    [10, "SmallRig VB99 Pro Mini V-Mount Battery", 249.99, None, True, None],
    [None, None, None, None, False, None],
    ["G&E", None, None, None, None, None],
    [None, "80c", None, None, None, None],
    [3, "Aputure STORM 80c BLAIR-CG LED Monolight (US)", 599, None, True, None],
    [2, "Aputure CF4 Fresnel for STORM 80c", 109, None, True, None],
    [2, "Aputure Barndoors for CF4 Fresnel", 45, None, True, None],
    [2, "Gator Stand Bag", 29.99, None, False, None],
    ["=$A$3", "Post", "=$C$3", None, "=$E$3", None],
    [1, "Apple MacBook Pro", 4099, None, True, None],
    [None, None, None, "Equipment Total", None, "Rental Cost"],
]


def test_reads_a_hand_made_list():
    got = importer.parse_rows(ROWS)
    names = [i.name for i in got.items]
    assert len(names) == 10 and "Equipment Total" not in names
    by = {i.name: i for i in got.items}
    assert by["Sony FX6 Full-Frame Cinema Camera (Body Only)"].department == "Camera"
    assert by["Sony FX6 Full-Frame Cinema Camera (Body Only)"].group == "Body"
    assert by["SmallRig VB99 Pro Mini V-Mount Battery"].qty == 10
    assert by["Aputure STORM 80c BLAIR-CG LED Monolight (US)"].group == "80c"
    assert by["Aputure STORM 80c BLAIR-CG LED Monolight (US)"].department == "G&E"
    assert by["Apple MacBook Pro"].department == "Post"
    assert not by["Gator Stand Bag"].rentable
    assert got.rate_pct == "4" and got.value_tax == "8.25"


def test_attaches_what_says_what_its_for():
    got = importer.parse_rows(ROWS)
    importer.attach_guesses(got.items)
    by = {i.name: i for i in got.items}
    fx6 = by["Sony FX6 Full-Frame Cinema Camera (Body Only)"]
    light = by["Aputure STORM 80c BLAIR-CG LED Monolight (US)"]
    assert by["Tilta Top Handle for Sony FX6"].parent == fx6.id
    assert by["Kondor Blue XLR Handle Extension for FX6 Top Handle"].parent == fx6.id  # not the handle
    assert by["Aputure CF4 Fresnel for STORM 80c"].parent == light.id
    assert by["Aputure Barndoors for CF4 Fresnel"].parent == light.id  # flattened to the light
    assert not by["SmallRig VB99 Pro Mini V-Mount Battery"].parent
    assert not fx6.parent and not light.parent


def test_gear_written_to_csv_reads_back_the_same(tmp_path):
    fx6 = Item(name="FX6", department="Camera", group="Body", qty=1, value="5998", serials=["123"])
    cage = Item(name="Cage for FX6", department="Camera", group="Body", value="279", parent=fx6.id)
    out = importer.write_csv([fx6, cage], {fx6.id: fx6, cage.id: cage}, tmp_path / "gear.csv")
    got = importer.read(out)
    by = {i.name: i for i in got.items}
    assert by["FX6"].serials == ["123"] and by["FX6"].value == "5998"
    assert by["Cage for FX6"].parent == by["FX6"].id


def test_an_xlsx_gear_list_is_read_by_its_headings(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for row in ROWS:
        ws.append(row)
    path = tmp_path / "list.xlsx"
    wb.save(path)
    got = importer.read(path)
    assert len(got.items) == 10 and got.attached >= 3


def test_links_on_names_come_in(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    for row in ROWS:
        wb.active.append(row)
    for r in wb.active.iter_rows():
        if r[1].value == "Sony FX6 Full-Frame Cinema Camera (Body Only)":
            r[1].hyperlink = "https://www.example.com/fx6"
    path = tmp_path / "linked.xlsx"
    wb.save(path)
    got = importer.read(path)
    fx6 = next(i for i in got.items if i.name.startswith("Sony FX6 Full-Frame"))
    assert fx6.link == "https://www.example.com/fx6"
    assert sum(1 for i in got.items if i.link) == 1


def test_a_google_sheet_is_read_as_a_spreadsheet_with_its_links(tmp_path, monkeypatch):
    import io
    import urllib.request

    import openpyxl

    assert importer.google_export_url("https://example.com/x") is None
    url = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/edit#gid=0"
    assert importer.google_export_url(url).endswith("/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/export?format=xlsx")

    # Google writes links as =HYPERLINK(...) formulas
    wb = openpyxl.Workbook()
    for row in ROWS:
        wb.active.append(row)
    for r in wb.active.iter_rows():
        if r[1].value == "Tilta Top Handle for Sony FX6":
            r[1].value = '=HYPERLINK("https://www.example.com/handle","Tilta Top Handle for Sony FX6")'
    buf = io.BytesIO()
    wb.save(buf)

    class Answer(io.BytesIO):
        def __init__(self, data, kind):
            super().__init__(data)
            self.headers = {"Content-Type": kind}

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    answer = {"data": b"<html>Sign in</html>", "kind": "text/html"}
    monkeypatch.setattr(urllib.request, "urlopen", lambda *a, **k: Answer(answer["data"], answer["kind"]))
    with pytest.raises(OSError, match="Anyone with the link"):
        importer.download_google_sheet(url, tmp_path)
    answer.update(data=buf.getvalue(), kind="application/vnd.openxmlformats")
    heard = []
    path = importer.download_google_sheet(url, tmp_path, lambda n, total: heard.append((n, total)))
    assert heard and heard[-1] == (len(buf.getvalue()), 0)  # how much has come, Google not saying how much
    import threading

    stopped = threading.Event()
    stopped.set()
    assert importer.download_google_sheet(url, tmp_path / "stopped", cancel=stopped) is None
    assert not (tmp_path / "stopped").exists()
    got = importer.read(path)
    handle = next(i for i in got.items if i.name == "Tilta Top Handle for Sony FX6")
    assert handle.link == "https://www.example.com/handle"


def test_a_csv_saved_by_excel_on_windows_keeps_its_dashes_and_accents(tmp_path):
    from bodega.core import importer

    path = tmp_path / "gear.csv"
    path.write_bytes("Item,Qty,Price\nCafé Lens – 50mm,1,900\n".encode("cp1252"))
    rows = importer._rows_csv(path)
    assert any("Café Lens – 50mm" in [str(c) for c in row] for row in rows)


def test_you_choose_what_a_blank_quantity_or_rent_means():
    from bodega.core import importer

    rows = [["Item", "Qty", "Price", "Rent?"], ["Sony FX6", 1, 5998, "Yes"], ["C-Stand", None, 120, None]]
    usual = importer.parse_rows(rows)
    stand = next(i for i in usual.items if i.name == "C-Stand")
    assert stand.qty == 1 and stand.rentable and usual.assumed_qty == 1 and usual.assumed_rent == 1
    strict = importer.parse_rows(rows, blank_rent=False, no_qty="skip")
    assert [i.name for i in strict.items] == ["Sony FX6"] and strict.skipped_qty == 1
    kept = importer.parse_rows(rows, blank_rent=False)
    assert not next(i for i in kept.items if i.name == "C-Stand").rentable


def test_the_sheets_own_total_checks_what_was_read():
    from decimal import Decimal

    from bodega.core import importer

    rows = [
        ["Item", "Qty", "Price"],
        ["Sony FX6", 1, 5998],
        ["Mini V-Mount", 6, 225],
        ["Equipment Total", None, 7348],
    ]
    got = importer.parse_rows(rows)
    assert got.sheet_total == Decimal(7348) and "matches" in importer.total_check(got)  # 5,998 + 6 × 225
    rows[-1][2] = 9000
    said = importer.total_check(importer.parse_rows(rows))
    assert "$9,000.00" in said and "$1,652.00 less" in said
    assert importer.total_check(importer.parse_rows(rows[:-1])) == ""  # no total row: nothing to check


def test_kits_and_rentals_export_as_csv_with_their_prices(tmp_path, data_folder):
    import csv

    from bodega.core import pricing, sample
    from bodega.core.store import Library

    lib = sample.fill(Library(data_folder).load())
    b = lib.pricing()
    kits = importer.write_kits_csv(lib, list(lib.packages.values()), b, tmp_path / "kits.csv")
    rows = list(csv.DictReader(kits.open(encoding="utf-8-sig")))
    p = next(iter(lib.packages.values()))
    assert rows[0]["Name"] == p.name and rows[0]["Day Rate"] == str(pricing.package_rate(lib, p, b))
    rentals = importer.write_rentals_csv(lib, list(lib.rentals.values()), tmp_path / "rentals.csv")
    got = list(csv.DictReader(rentals.open(encoding="utf-8-sig")))
    r = next(iter(lib.rentals.values()))
    assert got[0]["Rental"] == r.number
    assert got[0]["Total"] == str(pricing.totals(lib, r, lib.businesses.get(r.business_id)).total)
